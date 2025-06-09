import os
import re
import sys
import tkinter as tk
from tkinter import filedialog, scrolledtext, messagebox, ttk
from pathlib import Path
from datetime import datetime
import docx
from PyPDF2 import PdfReader
import traceback


def resource_path(relative_path):
    """获取打包后资源的绝对路径"""
    try:
        base_path = sys._MEIPASS
    except AttributeError:
        base_path = os.path.abspath(".")
    return os.path.join(base_path, relative_path)


class FileRenameApp:
    def __init__(self, root):
        self.root = root
        self.root.title("多功能文件重命名工具")
        self.root.geometry("900x750")  # 增加窗口高度
        self.font = ('Microsoft YaHei UI', 10)

        # 存储预览结果
        self.preview_results = []
        self.is_single_file = False
        self.current_mode = "title"

        # 创建界面元素
        self.create_widgets()
        print("[INFO] 界面初始化完成")

    def create_widgets(self):
        # 模式选择框架
        mode_frame = tk.Frame(self.root, padx=10, pady=5)
        mode_frame.pack(fill=tk.X)

        self.mode_var = tk.StringVar(value="title")
        batch_radio = tk.Radiobutton(mode_frame, text="批量重命名",
                                     variable=self.mode_var, value="batch",
                                     font=self.font, command=self.change_mode)
        batch_radio.pack(side=tk.LEFT, padx=10)

        single_radio = tk.Radiobutton(mode_frame, text="单个文件重命名",
                                      variable=self.mode_var, value="single",
                                      font=self.font, command=self.change_mode)
        single_radio.pack(side=tk.LEFT, padx=10)

        title_radio = tk.Radiobutton(mode_frame, text="按文件标题重命名",
                                     variable=self.mode_var, value="title",
                                     font=self.font, command=self.change_mode)
        title_radio.pack(side=tk.LEFT, padx=10)

        # 路径选择框架
        path_frame = tk.Frame(self.root, padx=10, pady=10)
        path_frame.pack(fill=tk.X)

        tk.Label(path_frame, text="文件/文件夹路径:", font=self.font).pack(side=tk.LEFT)

        self.path_entry = tk.Entry(path_frame, width=50, font=self.font)
        self.path_entry.pack(side=tk.LEFT, padx=5)

        self.folder_btn = tk.Button(path_frame, text="浏览文件夹", font=self.font,
                                    command=self.browse_folder)
        self.folder_btn.pack(side=tk.LEFT, padx=5)

        self.file_btn = tk.Button(path_frame, text="浏览文件", font=self.font,
                                  command=self.browse_file)
        self.file_btn.pack(side=tk.LEFT, padx=5)

        self.mode_label = tk.Label(path_frame, text="按标题重命名模式", font=self.font, fg="blue")
        self.mode_label.pack(side=tk.LEFT, padx=10)

        # 标题重命名选项框架
        self.title_options_frame = tk.Frame(self.root, padx=10, pady=5)
        self.title_options_frame.pack(fill=tk.X)

        self.auto_rename_title_var = tk.BooleanVar(value=False)  # 默认不自动重命名
        auto_rename_title_check = tk.Checkbutton(self.title_options_frame, text="自动重命名",
                                                 variable=self.auto_rename_title_var, font=self.font)
        auto_rename_title_check.pack(side=tk.LEFT, padx=5)

        self.title_prefix_var = tk.StringVar(value="")
        prefix_label = tk.Label(self.title_options_frame, text="前缀:", font=self.font)
        prefix_label.pack(side=tk.LEFT, padx=5)
        prefix_entry = tk.Entry(self.title_options_frame, textvariable=self.title_prefix_var, width=10, font=self.font)
        prefix_entry.pack(side=tk.LEFT, padx=5)

        self.title_suffix_var = tk.StringVar(value="")
        suffix_label = tk.Label(self.title_options_frame, text="后缀:", font=self.font)
        suffix_label.pack(side=tk.LEFT, padx=5)
        suffix_entry = tk.Entry(self.title_options_frame, textvariable=self.title_suffix_var, width=10, font=self.font)
        suffix_entry.pack(side=tk.LEFT, padx=5)

        # 新增：文件格式后缀修改框架
        self.extension_frame = tk.Frame(self.root, padx=10, pady=5)
        self.extension_frame.pack(fill=tk.X)

        tk.Label(self.extension_frame, text="修改文件后缀:", font=self.font).pack(side=tk.LEFT, padx=5)

        # 常用文件格式下拉菜单
        self.common_extensions = [
            "不修改", "txt", "pdf", "docx", "doc", "xlsx", "xls", "pptx", "ppt",
            "jpg", "jpeg", "png", "gif", "bmp", "svg", "mp4", "avi", "mov",
            "mp3", "wav", "zip", "rar", "7z", "html", "htm", "json", "csv"
        ]

        self.extension_var = tk.StringVar(value="不修改")
        self.extension_combobox = ttk.Combobox(
            self.extension_frame,
            textvariable=self.extension_var,
            values=self.common_extensions,
            width=10,
            font=self.font
        )
        self.extension_combobox.pack(side=tk.LEFT, padx=5)
        self.extension_combobox.bind("<<ComboboxSelected>>", self.on_extension_change)

        # 自定义后缀输入框
        tk.Label(self.extension_frame, text="自定义:", font=self.font).pack(side=tk.LEFT, padx=5)
        self.custom_ext_var = tk.StringVar(value="")
        self.custom_ext_entry = tk.Entry(
            self.extension_frame,
            textvariable=self.custom_ext_var,
            width=10,
            font=self.font
        )
        self.custom_ext_entry.pack(side=tk.LEFT, padx=5)
        self.custom_ext_entry.bind("<KeyRelease>", self.on_custom_ext_change)

        # 预览结果框架
        preview_frame = tk.Frame(self.root, padx=10, pady=10)
        preview_frame.pack(fill=tk.BOTH, expand=True)

        # 左侧 - 原始文件
        left_frame = tk.Frame(preview_frame)
        left_frame.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 5))

        self.original_label = tk.Label(left_frame, text="原始文件名:", font=self.font, fg="blue")
        self.original_label.pack(anchor=tk.W)

        self.original_files_text = scrolledtext.ScrolledText(left_frame, width=40, height=5, font=self.font)
        self.original_files_text.pack(fill=tk.BOTH, expand=True, pady=5)

        # 右侧 - 提取结果
        right_frame = tk.Frame(preview_frame)
        right_frame.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(5, 0))

        self.preview_label = tk.Label(right_frame, text="提取结果:", font=self.font, fg="green")
        self.preview_label.pack(anchor=tk.W)

        self.preview_text = scrolledtext.ScrolledText(right_frame, width=40, height=5, font=self.font)
        self.preview_text.pack(fill=tk.BOTH, expand=True, pady=5)

        # 按钮框架
        btn_frame = tk.Frame(self.root, padx=10, pady=10)
        btn_frame.pack(fill=tk.X)

        self.extract_title_btn = tk.Button(btn_frame, text="提取标题", font=self.font,
                                           command=self.extract_title)
        self.extract_title_btn.pack(side=tk.LEFT, padx=5)

        self.execute_btn = tk.Button(btn_frame, text="执行重命名", font=self.font,
                                     command=self.execute_rename, state=tk.DISABLED)
        self.execute_btn.pack(side=tk.LEFT, padx=5)

        # 状态框架
        status_frame = tk.Frame(self.root, padx=10, pady=5)
        status_frame.pack(fill=tk.X)

        self.status_var = tk.StringVar(value="就绪")
        status_label = tk.Label(status_frame, textvariable=self.status_var, font=self.font, fg="blue")
        status_label.pack(anchor=tk.W)

        # 初始化界面
        self.change_mode()
        print("[INFO] 界面组件加载完成")

    def change_mode(self):
        self.current_mode = self.mode_var.get()
        print(f"[INFO] 切换到{self.current_mode}模式")

        if self.current_mode == "batch":
            self.is_single_file = False
            self.mode_label.config(text="批量处理模式")
            self.folder_btn.config(text="浏览文件夹", command=self.browse_folder)
            self.title_options_frame.pack_forget()
            self.extract_title_btn.pack_forget()
            self.original_label.config(text="原始文件列表:")
            self.preview_label.config(text="重命名预览:")

        elif self.current_mode == "single":
            self.is_single_file = True
            self.mode_label.config(text="单个文件模式")
            self.folder_btn.config(text="浏览文件", command=self.browse_file)
            self.title_options_frame.pack_forget()
            self.extract_title_btn.pack_forget()
            self.original_label.config(text="原始文件:")
            self.preview_label.config(text="重命名预览:")

        elif self.current_mode == "title":
            self.is_single_file = True
            self.mode_label.config(text="按标题重命名模式")
            self.folder_btn.config(text="浏览文件", command=self.browse_file)
            self.title_options_frame.pack(fill=tk.X)
            self.extract_title_btn.pack(side=tk.LEFT, padx=5)
            self.original_label.config(text="原始文件名:")
            self.preview_label.config(text="提取结果:")

        self.path_entry.delete(0, tk.END)
        self.original_files_text.delete(1.0, tk.END)
        self.preview_text.delete(1.0, tk.END)
        self.execute_btn.config(state=tk.DISABLED)
        self.status_var.set("就绪")
        print("[INFO] 模式切换完成")

    def browse_folder(self):
        print("[INFO] 打开文件/文件夹选择对话框")
        if self.current_mode == "batch":
            folder_selected = filedialog.askdirectory()
            if folder_selected:
                self.path_entry.delete(0, tk.END)
                self.path_entry.insert(0, folder_selected)
                self.status_var.set(f"已选择文件夹: {os.path.basename(folder_selected)}")
                print(f"[INFO] 选择文件夹: {folder_selected}")

                # 批量模式下自动显示文件列表
                self.show_batch_files(folder_selected)
        else:  # single 或 title 模式
            file_selected = filedialog.askopenfilename(
                filetypes=[("所有文件", "*.*")]
            )
            if file_selected:
                self.path_entry.delete(0, tk.END)
                self.path_entry.insert(0, file_selected)
                self.status_var.set(f"已选择文件: {os.path.basename(file_selected)}")
                print(f"[INFO] 选择文件: {file_selected}")

    def browse_file(self):
        print("[INFO] 打开文件选择对话框")
        if self.current_mode == "title":
            file_selected = filedialog.askopenfilename(
                filetypes=[("Word/PDF文件", "*.docx *.doc *.pdf"),
                           ("Word文件", "*.docx *.doc"),
                           ("PDF文件", "*.pdf"),
                           ("所有文件", "*.*")]
            )
        else:
            file_selected = filedialog.askopenfilename(
                filetypes=[("所有文件", "*.*")]
            )

        if file_selected:
            self.path_entry.delete(0, tk.END)
            self.path_entry.insert(0, file_selected)
            self.status_var.set(f"已选择文件: {os.path.basename(file_selected)}")
            print(f"[INFO] 选择文件: {file_selected}")

    def show_batch_files(self, folder_path):
        """显示批量模式下的文件列表"""
        self.original_files_text.delete(1.0, tk.END)
        try:
            files = os.listdir(folder_path)
            files = [f for f in files if os.path.isfile(os.path.join(folder_path, f))]

            if not files:
                self.original_files_text.insert(tk.END, "所选文件夹为空")
                return

            for file in files[:100]:  # 限制显示数量
                self.original_files_text.insert(tk.END, file + "\n")

            if len(files) > 100:
                self.original_files_text.insert(tk.END, f"... 共{len(files)}个文件")

            # 自动生成预览
            self.generate_batch_preview(files, folder_path)

        except Exception as e:
            self.original_files_text.insert(tk.END, f"无法读取文件夹内容: {str(e)}")

    def generate_batch_preview(self, files, folder_path):
        """生成批量重命名预览"""
        self.preview_text.delete(1.0, tk.END)

        prefix = self.title_prefix_var.get().strip()
        suffix = self.title_suffix_var.get().strip()
        target_ext = self.get_target_extension()

        renamed_files = []

        for file in files:
            base_name, ext = os.path.splitext(file)

            # 应用前缀和后缀
            new_base_name = f"{prefix}{base_name}{suffix}"

            # 应用新的文件后缀
            if target_ext:
                new_ext = f".{target_ext}"
            else:
                new_ext = ext

            new_name = new_base_name + new_ext
            renamed_files.append((file, new_name))

            # 只显示前100个预览
            if len(renamed_files) <= 100:
                self.preview_text.insert(tk.END, f"{file} → {new_name}\n")

        if len(files) > 100:
            self.preview_text.insert(tk.END, f"... 共{len(files)}个文件将被重命名")

        self.preview_results = renamed_files
        self.execute_btn.config(state=tk.NORMAL if renamed_files else tk.DISABLED)
        print(f"[INFO] 生成批量重命名预览，{len(renamed_files)}个文件")

    def on_extension_change(self, event=None):
        """下拉菜单选择变化时的处理"""
        selected = self.extension_var.get()
        if selected != "不修改":
            self.custom_ext_var.set("")  # 清空自定义输入框

        # 如果是批量模式，自动更新预览
        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                files = os.listdir(folder_path)
                files = [f for f in files if os.path.isfile(os.path.join(folder_path, f))]
                self.generate_batch_preview(files, folder_path)

    def on_custom_ext_change(self, event=None):
        """自定义后缀输入变化时的处理"""
        custom_ext = self.custom_ext_var.get().strip()
        if custom_ext:
            self.extension_var.set("不修改")  # 取消下拉菜单选择

        # 如果是批量模式，自动更新预览
        if self.current_mode == "batch" and self.path_entry.get():
            folder_path = self.path_entry.get()
            if os.path.isdir(folder_path):
                files = os.listdir(folder_path)
                files = [f for f in files if os.path.isfile(os.path.join(folder_path, f))]
                self.generate_batch_preview(files, folder_path)

    def get_target_extension(self):
        """获取目标文件后缀"""
        selected_ext = self.extension_var.get()
        custom_ext = self.custom_ext_var.get().strip()

        if selected_ext != "不修改":
            return selected_ext
        elif custom_ext:
            # 清理自定义后缀，移除可能的点号
            return custom_ext.replace('.', '')
        else:
            return None

    def extract_title_from_docx(self, file_path):
        """从Word文档中提取标题"""
        try:
            print(f"[INFO] 开始从Word文档提取标题: {file_path}")
            doc = docx.Document(file_path)

            # 方法1：优先从文档属性提取
            if doc.core_properties.title:
                title = doc.core_properties.title.strip()
                print(f"[INFO] 从文档属性提取标题: {title}")
                return title

            # 方法2：针对"中华人民共和国劳动合同法"文档的强化提取
            for para in doc.paragraphs[:10]:
                text = para.text.strip()
                if ("中华人民共和国" in text or "法" in text) and len(text) < 60:
                    if not re.match(r'^(第[一二三四五六七八九十0-9]+章|目?录|附件|附录|修订说明|前言|序)', text):
                        print(f"[INFO] 从正文提取标题: {text}")
                        return text[:40]

            # 方法3：常规标题提取逻辑
            title_candidates = []

            for para in doc.paragraphs:
                if para.style.name.startswith('Heading') or para.style.name.lower().startswith('标题'):
                    candidate = para.text.strip()
                    if len(candidate) > 5 and not re.match(r'^(第[0-9一二三四五六七八九十]+)', candidate):
                        title_candidates.append(candidate)
                        print(f"[INFO] 从标题样式提取: {candidate}")
                        break

            if not title_candidates:
                for para in doc.paragraphs[:3]:
                    text = para.text.strip()
                    if len(text) > 15 and not re.match(r'^(第[0-9一二三四五六七八九十]+|目?录|附件|附录)', text):
                        title_candidates.append(text[:40])
                        print(f"[INFO] 从正文前3段提取: {text[:40]}")
                        break

            if not title_candidates and doc.paragraphs:
                first_line = doc.paragraphs[0].text.strip()
                if len(first_line) > 5:
                    title_candidates.append(first_line[:40])
                    print(f"[INFO] 从第一行提取: {first_line[:40]}")

            return title_candidates[0] if title_candidates else "无标题"

        except Exception as e:
            error_msg = f"提取失败: {str(e)}"
            print(f"[ERROR] Word标题提取失败: {error_msg}")
            return error_msg

    def extract_title_from_pdf(self, file_path):
        """从PDF文件中提取标题"""
        try:
            print(f"[INFO] 开始从PDF提取标题: {file_path}")
            pdf = PdfReader(file_path)
            info = pdf.metadata
            if info and '/Title' in info:
                title = info['/Title']
                if isinstance(title, bytes):
                    title = title.decode('utf-8', errors='replace')
                print(f"[INFO] 从PDF元数据提取标题: {title}")
                return title.strip()

            if len(pdf.pages) > 0:
                first_page = pdf.pages[0].extract_text()
                if not first_page:
                    print("[INFO] PDF首页无文本内容")
                    return "无标题"

                lines = []
                for line in first_page.split('\n'):
                    stripped = line.strip()
                    if stripped and len(stripped) > 5:
                        lines.append(stripped)

                if lines:
                    longest_line = max(lines, key=len)
                    print(f"[INFO] 从PDF正文提取标题: {longest_line[:40]}")
                    return longest_line[:40]
                else:
                    print("[INFO] PDF首页无有效行")
                    return "无标题"

            print("[INFO] PDF无有效内容，返回无标题")
            return "无标题"
        except Exception as e:
            error_msg = f"提取失败: {str(e)}"
            print(f"[ERROR] PDF标题提取失败: {error_msg}")
            return error_msg

    def extract_title(self):
        file_path = self.path_entry.get().strip()
        if not file_path:
            messagebox.showerror("错误", "请选择文件")
            print("[ERROR] 未选择文件")
            return

        if not os.path.exists(file_path):
            messagebox.showerror("错误", "文件不存在")
            print(f"[ERROR] 文件不存在: {file_path}")
            return

        file_ext = os.path.splitext(file_path)[1].lower()
        filename = os.path.basename(file_path)

        self.original_files_text.delete(1.0, tk.END)
        self.original_files_text.insert(tk.END, filename)

        self.preview_text.delete(1.0, tk.END)

        try:
            self.status_var.set(f"正在提取标题: {filename}")
            print(f"[INFO] 开始提取标题流程: {filename}")

            if file_ext in ['.docx', '.doc']:
                title = self.extract_title_from_docx(file_path)
            elif file_ext == '.pdf':
                title = self.extract_title_from_pdf(file_path)
            else:
                title = f"不支持的文件格式: {file_ext}"
                print(f"[ERROR] 不支持的文件格式: {file_ext}")

            if title.startswith("提取失败"):
                self.preview_text.insert(tk.END, f"错误: {title}\n", "error")
                self.status_var.set("提取失败")
                self.execute_btn.config(state=tk.DISABLED)
                print(f"[ERROR] 标题提取失败: {title}")
                return

            valid_title = re.sub(r'[\\/:*?"<>|]', '_', title)
            if not valid_title:
                valid_title = "无标题_" + datetime.now().strftime("%Y%m%d%H%M%S")
                print("[INFO] 生成默认标题: 无标题_时间戳")

            prefix = self.title_prefix_var.get().strip()
            suffix = self.title_suffix_var.get().strip()

            # 应用文件后缀修改
            target_ext = self.get_target_extension()
            if target_ext:
                new_ext = f".{target_ext}"
            else:
                new_ext = os.path.splitext(filename)[1]

            new_filename = f"{prefix}{valid_title}{suffix}{new_ext}"

            self.preview_text.insert(tk.END, f"原始文件名: {filename}\n\n", "original")
            self.preview_text.insert(tk.END, f"提取的标题: {title}\n\n", "title")
            self.preview_text.insert(tk.END, f"新文件名: {new_filename}\n", "new")

            self.status_var.set(f"标题提取完成: {title}")
            self.preview_results = [(filename, new_filename)]
            self.file_path = file_path

            self.execute_btn.config(state=tk.NORMAL)
            print(f"[INFO] 标题提取成功，新文件名: {new_filename}")

            if self.auto_rename_title_var.get():
                print("[INFO] 自动重命名已启用，执行重命名")
                self.execute_rename()

        except Exception as e:
            self.preview_text.insert(tk.END, f"错误: {str(e)}\n", "error")
            self.status_var.set(f"提取出错: {str(e)}")
            traceback.print_exc()
            self.execute_btn.config(state=tk.DISABLED)
            print(f"[ERROR] 提取标题过程中出错: {str(e)}")

    def execute_rename(self):
        if not self.preview_results:
            messagebox.showinfo("提示", "没有可执行的重命名操作")
            print("[INFO] 没有可执行的重命名操作")
            return

        try:
            if self.current_mode == "batch":
                folder_path = self.path_entry.get()
                if not os.path.isdir(folder_path):
                    messagebox.showerror("错误", "请选择有效的文件夹路径")
                    return

                renamed_count = 0
                failed_files = []

                for old_name, new_name in self.preview_results:
                    old_path = os.path.join(folder_path, old_name)
                    new_path = os.path.join(folder_path, new_name)

                    try:
                        os.rename(old_path, new_path)
                        renamed_count += 1
                    except Exception as e:
                        failed_files.append(f"{old_name} → {new_name}: {str(e)}")

                success_msg = f"成功重命名 {renamed_count} 个文件"
                if failed_files:
                    success_msg += f"\n\n{len(failed_files)} 个文件重命名失败:"
                    for fail in failed_files[:10]:  # 只显示前10个失败
                        success_msg += f"\n{fail}"
                    if len(failed_files) > 10:
                        success_msg += f"\n... 等{len(failed_files)}个文件"

                messagebox.showinfo("成功", success_msg)
                self.status_var.set(f"批量重命名完成: {renamed_count}个文件成功")
                print(f"[INFO] 批量重命名完成，成功: {renamed_count}，失败: {len(failed_files)}")

                # 刷新文件列表
                self.show_batch_files(folder_path)

            else:  # 单文件模式
                old_name, new_name = self.preview_results[0]
                old_path = self.file_path
                new_path = os.path.join(os.path.dirname(old_path), new_name)

                os.rename(old_path, new_path)
                self.preview_text.insert(tk.END, f"\n\n重命名成功: {old_name} → {new_name}\n", "success")

                self.status_var.set(f"重命名完成: {new_name}")
                messagebox.showinfo("成功", f"已成功重命名文件:\n{old_name} → {new_name}")

                self.path_entry.delete(0, tk.END)
                self.path_entry.insert(0, new_path)
                self.original_files_text.delete(1.0, tk.END)
                self.original_files_text.insert(tk.END, new_name)

                print(f"[INFO] 文件重命名成功: {old_path} → {new_path}")

        except Exception as e:
            messagebox.showerror("错误", f"执行时出错: {str(e)}")
            self.status_var.set("操作失败")
            print(f"[ERROR] 文件重命名失败: {str(e)}")


if __name__ == "__main__":
    print("[INFO] 程序启动中...")
    root = tk.Tk()
    app = FileRenameApp(root)
    print("[INFO] 程序初始化完成，进入主事件循环")
    root.mainloop()